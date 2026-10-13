"""Summaries and spreadsheet export, grouped by date, place, reason, vendor..."""

from __future__ import annotations

import csv
from collections import defaultdict
from pathlib import Path
from typing import Callable, Iterable, Optional

from .db import Database
from .models import COUNTED_STATUSES, DUPLICATE, EXCLUDED, NEEDS_REVIEW


def _date(r) -> str:
    return (r["purchase_date"] or (r["email_date"] or "")[:10] or "")


def _place(r) -> str:
    return ", ".join(p for p in (r["city"], r["region"]) if p) or (r["country"] or "Unknown")


GROUPERS: dict[str, Callable] = {
    "month": lambda r: _date(r)[:7] or "Unknown",
    "quarter": lambda r: f"{_date(r)[:4]}-Q{(int(_date(r)[5:7]) - 1) // 3 + 1}" if len(_date(r)) >= 7 else "Unknown",
    "year": lambda r: _date(r)[:4] or "Unknown",
    "category": lambda r: r["category"] or "Uncategorized",
    "place": _place,
    "city": lambda r: r["city"] or "Unknown",
    "region": lambda r: r["region"] or "Unknown",
    "country": lambda r: r["country"] or "Unknown",
    "vendor": lambda r: r["vendor"] or "Unknown",
    "project": lambda r: r["project"] or "No job listed",
    "account": lambda r: r["account"],
}


def paid_purchases(db: Database, since: Optional[str] = None, until: Optional[str] = None,
                   category: Optional[str] = None, place: Optional[str] = None) -> list:
    rows = [r for r in db.records(COUNTED_STATUSES) if r["paid_status"] == "paid"]
    if since:
        rows = [r for r in rows if _date(r) >= since]
    if until:
        rows = [r for r in rows if _date(r) and _date(r) <= until]
    if category:
        rows = [r for r in rows if (r["category"] or "").lower() == category.lower()]
    if place:
        p = place.lower()
        rows = [r for r in rows if p in " ".join(filter(None, (r["city"], r["region"], r["country"], r["address"]))).lower()]
    return rows


def group(rows: Iterable, by: str) -> list[tuple[str, int, dict[str, float]]]:
    key = GROUPERS[by]
    buckets: dict[str, list] = defaultdict(list)
    for r in rows:
        buckets[key(r)].append(r)
    out = []
    for k in sorted(buckets):
        totals: dict[str, float] = defaultdict(float)
        for r in buckets[k]:
            if r["amount"] is not None:
                totals[r["currency"] or "?"] += r["amount"]
        out.append((k, len(buckets[k]), dict(totals)))
    if by not in ("month", "quarter", "year"):
        out.sort(key=lambda t: -sum(t[2].values()))
    return out


def _money(totals: dict[str, float]) -> str:
    return "  ".join(f"{v:,.2f} {c}" for c, v in sorted(totals.items())) or "-"


def print_summary(rows: list, by: str, out=print) -> None:
    groups = group(rows, by)
    width = max([len(by)] + [len(g[0]) for g in groups]) + 2
    out(f"{by.capitalize():<{width}}{'Count':>6}  Total")
    out("-" * (width + 30))
    for name, count, totals in groups:
        out(f"{name:<{width}}{count:>6}  {_money(totals)}")
    out("-" * (width + 30))
    all_totals: dict[str, float] = defaultdict(float)
    for _, _, t in groups:
        for c, v in t.items():
            all_totals[c] += v
    out(f"{'All':<{width}}{len(rows):>6}  {_money(all_totals)}")
    missing = sum(1 for r in rows if r["amount"] is None)
    if missing:
        out(f"({missing} paid item(s) have no amount - run `review --all` or `set` to add it)")


def print_list(rows: list, out=print) -> None:
    for r in rows:
        amt = f"{r['amount']:>10,.2f} {r['currency'] or ''}" if r["amount"] is not None else f"{'?':>10}    "
        out(f"#{r['email_id']:<5} {_date(r):<10} {amt}  {(r['vendor'] or '?')[:24]:<24} "
            f"{(r['category'] or '')[:24]:<24} {_place(r)[:20]:<20} {(r['purpose'] or '')[:40]}")


COLUMNS = [
    ("ID", "email_id"), ("Date paid", None), ("Vendor", "vendor"), ("Amount", "amount"), ("Currency", "currency"),
    ("Category", "category"), ("What for", "purpose"), ("Job / project", "project"), ("City", "city"), ("State/Region", "region"),
    ("Country", "country"), ("Address", "address"), ("Invoice/Order #", "invoice_number"),
    ("Type", "doc_type"), ("Proof of payment", "paid_evidence"), ("How settled", "status"),
    ("Confidence", "confidence"), ("Verification notes", "verify_note"), ("Email account", "account"),
    ("Email subject", "subject"), ("From", "from_addr"), ("Email date", "email_date"),
]


def _row_values(r) -> list:
    return [(_date(r) if key is None else r[key]) for _, key in COLUMNS]


def export_csv(rows: list, path: Path) -> None:
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow([c for c, _ in COLUMNS])
        for r in rows:
            w.writerow(_row_values(r))


def export_xlsx(db: Database, rows: list, path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    bold = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="DDE7F0")
    wb = Workbook()

    def sheet(title: str, header: list[str], data: Iterable[list], widths: Optional[list[int]] = None):
        ws = wb.create_sheet(title)
        ws.append(header)
        for cell in ws[1]:
            cell.font, cell.fill = bold, head_fill
        for row in data:
            ws.append(row)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for i, h in enumerate(header, 1):
            ws.column_dimensions[get_column_letter(i)].width = (widths[i - 1] if widths else max(12, min(40, len(h) + 4)))
        return ws

    wb.remove(wb.active)
    widths = [6, 11, 24, 11, 8, 26, 34, 20, 14, 12, 10, 24, 14, 18, 40, 14, 10, 50, 14, 40, 30, 22]
    ws = sheet("Paid purchases", [c for c, _ in COLUMNS], (_row_values(r) for r in rows), widths)
    for cell in ws["D"][1:]:
        cell.number_format = "#,##0.00"

    for by, title in (("month", "By month"), ("quarter", "By quarter"), ("category", "By category (reason)"), ("project", "By job or project"),
                      ("place", "By place"), ("vendor", "By vendor"), ("account", "By email account")):
        currencies = sorted({r["currency"] or "?" for r in rows if r["amount"] is not None}) or ["USD"]
        data = [[name, count, *[round(t.get(c, 0.0), 2) for c in currencies]] for name, count, t in group(rows, by)]
        sheet(title, [by.capitalize(), "Count", *[f"Total {c}" for c in currencies]], data, [30, 8] + [14] * len(currencies))

    # Month x category grid (main currency only).
    if rows:
        main_cur = max({r["currency"] or "?" for r in rows}, key=lambda c: sum(1 for r in rows if (r["currency"] or "?") == c))
        months = sorted({GROUPERS["month"](r) for r in rows})
        cats = sorted({GROUPERS["category"](r) for r in rows})
        grid = defaultdict(float)
        for r in rows:
            if r["amount"] is not None and (r["currency"] or "?") == main_cur:
                grid[(GROUPERS["month"](r), GROUPERS["category"](r))] += r["amount"]
        sheet(f"Month x category ({main_cur})", ["Month", *cats],
              ([m, *[round(grid[(m, c)], 2) or None for c in cats]] for m in months), [10] + [16] * len(cats))

    review_rows = db.records([NEEDS_REVIEW])
    sheet("Needs your review", ["ID", "Email date", "From", "Subject", "Best guess", "Amount", "Why"],
          ([r["email_id"], r["email_date"], r["from_addr"], r["subject"], r["paid_status"], r["amount"],
            r["verify_note"] or r["notes"]] for r in review_rows), [6, 22, 30, 40, 12, 11, 60])

    skipped = db.records([EXCLUDED, DUPLICATE])
    sheet("Not counted", ["ID", "Email date", "From", "Subject", "Type", "Reason", "Amount", "Note"],
          ([r["email_id"], r["email_date"], r["from_addr"], r["subject"], r["doc_type"],
            "Same purchase as #%s" % r["duplicate_of"] if r["status"] == DUPLICATE else "Not paid / not a purchase",
            r["amount"], r["verify_note"] or r["paid_evidence"]] for r in skipped),
          [6, 22, 30, 40, 18, 26, 11, 60])
    wb.save(path)
