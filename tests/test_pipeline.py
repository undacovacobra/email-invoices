from types import SimpleNamespace

from openpyxl import load_workbook

from invoice_sorter import prefilter, report
from invoice_sorter.analyzers import ClaudeAnalyzer, OfflineAnalyzer
from invoice_sorter.models import Extraction, Location, Verification
from invoice_sorter.parsing import parse_email
from invoice_sorter.pipeline import all_categories, scan
from invoice_sorter.review import examples_for_prompt, review, set_fields

from .conftest import SAMPLES, make_email


def by_subject(db):
    return {r["subject"]: r for r in db.records()}


def ex(**kw):
    base = dict(doc_type="receipt", paid_status="paid", paid_evidence="", vendor=None, amount=None,
                currency="USD", purchase_date=None, invoice_number=None,
                location=Location(city=None, region=None, country=None, address=None),
                purpose=None, project=None, category="Uncategorized", confidence=0.9, notes="")
    base.update(kw)
    return Extraction(**base)


class ScriptedAnalyzer:
    """Stands in for Claude: answers are keyed by subject."""

    name = "scripted"

    def __init__(self, answers, verdicts):
        self.answers, self.verdicts, self.verify_calls = answers, verdicts, []

    def extract(self, email):
        return self.answers[email.subject]

    def verify(self, email, prior, related):
        self.verify_calls.append((email.subject, [r.subject for _, r in related]))
        return self.verdicts[email.subject]


def test_parse_html_and_attachments():
    raw = make_email("Receipt", "plain", "A <a@x.com>", "Mon, 02 Mar 2026 10:00:00 +0000", "<x@x>",
                     html="<html><style>p{}</style><p>Total <b>$5.00</b></p></html>",
                     attachments=[("r.png", "image/png", b"\x89PNG..")])
    e = parse_email(raw, "acct")
    assert e.body_text == "plain"
    assert e.attachments[0].filename == "r.png"
    assert e.date.startswith("2026-03-02")


def test_prefilter_skips_newsletters():
    news = parse_email(SAMPLES["newsletter"], "a")
    receipt = parse_email(SAMPLES["receipt"], "a")
    assert prefilter.score(news.subject, news.body_text, []) < prefilter.MIN_SCORE
    assert prefilter.score(receipt.subject, receipt.body_text, []) >= prefilter.MIN_SCORE


def test_scan_with_ai_confirms_verifies_and_links(cfg, db):
    abc = dict(vendor="ABC Electric", currency="USD",
               location=Location(city="Austin", region="TX", country="US", address="12 Oak St"),
               purpose="Panel upgrade", category="Contractors & Professional Services")
    answers = {
        "Your Staples receipt": ex(vendor="Staples", amount=84.5, purchase_date="2026-03-02",
                                   location=Location(city="Austin", region="TX", country="US", address=None),
                                   purpose="Printer paper and toner", category="Office Supplies", confidence=0.97),
        "Quote Q-1001 for panel upgrade": ex(doc_type="quote_or_estimate", paid_status="not_paid", amount=2400.0,
                                             invoice_number="Q-1001", confidence=0.95, **abc),
        "Invoice INV-2001": ex(doc_type="unpaid_invoice", paid_status="not_paid", amount=2400.0,
                               invoice_number="INV-2001", purchase_date="2026-04-01", confidence=0.9, **abc),
        # Short acknowledgement: paid, but no details and not confident.
        "Re: Invoice INV-2001": ex(doc_type="payment_confirmation", paid_status="paid", vendor="ABC Electric",
                                   purchase_date="2026-04-10", confidence=0.6),
    }
    verdicts = {
        "Re: Invoice INV-2001": Verification(
            paid_status="paid", category="Contractors & Professional Services", confidence=0.92,
            reasoning="Reply in the invoice thread says payment received.", supporting_email_ids=[3],
            amount=None, purchase_date=None, vendor=None, purpose=None, project=None),
    }
    analyzer = ScriptedAnalyzer(answers, verdicts)
    result = scan(cfg, db, analyzer, log=lambda *_: None)

    recs = by_subject(db)
    assert "Spring newsletter" not in recs  # filtered before any AI call
    assert recs["Your Staples receipt"]["status"] == "confirmed"
    ack = recs["Re: Invoice INV-2001"]
    assert ack["status"] == "verified"
    assert ack["amount"] == 2400.0  # filled in from the matched invoice
    assert ack["invoice_number"] == "INV-2001"
    assert ack["city"] == "Austin"
    assert recs["Invoice INV-2001"]["status"] == "duplicate"
    assert recs["Invoice INV-2001"]["duplicate_of"] == ack["email_id"]
    # The quote for the same job (same vendor, same amount) is linked to the payment, not counted.
    quote = recs["Quote Q-1001 for panel upgrade"]
    assert (quote["status"], quote["duplicate_of"]) == ("duplicate", ack["email_id"])

    # The double-check saw the invoice it was replying to and the matching quote.
    assert analyzer.verify_calls == [("Re: Invoice INV-2001", ["Invoice INV-2001", "Quote Q-1001 for panel upgrade"])]

    paid = report.paid_purchases(db)
    assert sorted(r["amount"] for r in paid) == [84.5, 2400.0]
    groups = {g[0]: g[2] for g in report.group(paid, "category")}
    assert groups["Office Supplies"] == {"USD": 84.5}
    assert "needs_review" not in result


def test_uncertain_goes_to_review_and_answers_are_learned(cfg, db):
    answers = {s: ex(paid_status="unclear", confidence=0.3) for s in
               ("Your Staples receipt", "Quote Q-1001 for panel upgrade", "Invoice INV-2001", "Re: Invoice INV-2001")}
    verdicts = {s: Verification(paid_status="unclear", category="Uncategorized", confidence=0.4,
                                reasoning="Can't tell.", supporting_email_ids=[], amount=None,
                                purchase_date=None, vendor=None, purpose=None, project=None) for s in answers}
    scan(cfg, db, ScriptedAnalyzer(answers, verdicts), log=lambda *_: None)
    assert {r["status"] for r in db.records()} == {"needs_review"}

    cats = all_categories(cfg, db)
    # Items are reviewed in date order: receipt, quote, invoice, ack.
    replies = iter([
        "y", "1", "84.50", "", "", "Staples", "Austin", "TX", "Paper", "",   # receipt -> Office Supplies
        "n",                                                             # quote
        "s",                                                             # invoice: skip
        "q",
    ])
    out = []
    done = review(db, cats, ask=lambda _: next(replies), out=out.append)
    assert done == 2
    recs = by_subject(db)
    r = recs["Your Staples receipt"]
    assert (r["status"], r["paid_status"], r["category"], r["amount"]) == ("user_confirmed", "paid", "Office Supplies", 84.5)
    assert recs["Quote Q-1001 for panel upgrade"]["status"] == "excluded"
    assert recs["Invoice INV-2001"]["status"] == "needs_review"

    examples = examples_for_prompt(db)
    assert {e["subject"] for e in examples} == {"Your Staples receipt", "Quote Q-1001 for panel upgrade"}

    # A new category typed by hand becomes available next time.
    set_fields(db, recs["Re: Invoice INV-2001"]["email_id"], cats, paid=True, category="Building Upgrades",
               amount=2400.0)
    assert "Building Upgrades" in all_categories(cfg, db)


def test_offline_mode_links_reply_to_invoice_but_asks_you(cfg, db):
    scan(cfg, db, OfflineAnalyzer(cfg.categories), log=lambda *_: None)
    recs = by_subject(db)
    ack = recs["Re: Invoice INV-2001"]
    assert ack["paid_status"] == "paid"
    assert ack["amount"] == 2400.0
    assert ack["status"] == "needs_review"  # offline mode never decides on its own
    assert recs["Invoice INV-2001"]["status"] == "duplicate"


def test_export_xlsx(cfg, db, tmp_path):
    scan(cfg, db, OfflineAnalyzer(cfg.categories), log=lambda *_: None)
    for r in db.records(["needs_review"]):
        set_fields(db, r["email_id"], cfg.categories, paid=True, category="Other")
    out = tmp_path / "out.xlsx"
    report.export_xlsx(db, report.paid_purchases(db), out)
    wb = load_workbook(out)
    assert {"Paid purchases", "By month", "By category (reason)", "By place", "Needs your review",
            "Not counted"} <= set(wb.sheetnames)
    assert wb["Paid purchases"].max_row >= 3


class FakeMessages:
    def __init__(self, parsed, stop_reason="end_turn"):
        self.parsed, self.stop_reason, self.calls = parsed, stop_reason, []

    def parse(self, **kwargs):
        self.calls.append(kwargs)
        return SimpleNamespace(parsed_output=self.parsed, stop_reason=self.stop_reason)


def test_claude_analyzer_request_shape():
    parsed = ex(category="Made Up Category", confidence=1.4, currency="usd")
    messages = FakeMessages(parsed)
    client = SimpleNamespace(beta=SimpleNamespace(messages=messages))
    analyzer = ClaudeAnalyzer("claude-opus-5-5", ["Office Supplies"], client=client,
                              examples=[{"from_addr": "a@b", "subject": "s", "paid_status": "paid",
                                         "category": "Office Supplies", "purpose": None}])
    raw = make_email("Receipt", "Total $5.00 paid", "A <a@x.com>", "Mon, 02 Mar 2026 10:00:00 +0000", "<x@x>",
                     attachments=[("receipt.pdf", "application/pdf", b"%PDF-1.4 fake")])
    result = analyzer.extract(parse_email(raw, "acct"))

    call = messages.calls[0]
    assert call["model"] == "claude-opus-5-5"
    assert call["fallbacks"] == "default" and call["betas"] == ["server-side-fallback-2026-07-01"]
    assert call["output_format"] is Extraction
    content = call["messages"][0]["content"]
    assert content[0]["type"] == "document" and content[-1]["type"] == "text"
    assert "Office Supplies" in call["system"] and "Decisions the user made" in call["system"]
    # Unknown categories are never invented; confidence is clamped.
    assert result.category == "Uncategorized" and "Made Up Category" in result.notes
    assert result.confidence == 1.0 and result.currency == "USD"


def test_claude_refusal_goes_to_review(cfg, db):
    messages = FakeMessages(None, stop_reason="refusal")
    client = SimpleNamespace(beta=SimpleNamespace(messages=messages))
    scan(cfg, db, ClaudeAnalyzer("claude-opus-5-5", cfg.categories, client=client), log=lambda *_: None)
    assert {r["status"] for r in db.records()} == {"needs_review"}


def test_rejecting_a_payment_brings_back_what_was_linked_to_it(cfg, db):
    scan(cfg, db, OfflineAnalyzer(cfg.categories), log=lambda *_: None)
    ack = by_subject(db)["Re: Invoice INV-2001"]
    set_fields(db, ack["email_id"], cfg.categories, paid=False)
    recs = by_subject(db)
    assert recs["Invoice INV-2001"]["status"] == "needs_review"
    assert recs["Invoice INV-2001"]["duplicate_of"] is None


def test_review_stops_cleanly_when_input_ends(cfg, db):
    scan(cfg, db, OfflineAnalyzer(cfg.categories), log=lambda *_: None)

    def ask(_):
        raise EOFError

    out = []
    review(db, cfg.categories, ask=ask, out=out.append)
    assert "Stopped" in out[-1]
